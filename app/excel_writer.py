"""
excel_writer.py
----------------
Builds the final QnA.xlsx in memory (as bytes) with three sheets —
English, Hindi, Marathi — each with Questions / Answers columns.
Returning bytes (instead of writing straight to disk) makes this
work cleanly with Streamlit's download_button.
"""

import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

SHEET_ORDER = ["English", "Hindi", "Marathi"]
HEADER_COLOR = "7A2E2E"  # TriBhasha maroon


class ExcelBuildError(Exception):
    pass


def build_excel(english_pairs: list, hindi_pairs: list, marathi_pairs: list) -> bytes:
    sheets = {
        "English": english_pairs,
        "Hindi": hindi_pairs,
        "Marathi": marathi_pairs,
    }

    for name in SHEET_ORDER:
        if not sheets[name]:
            raise ExcelBuildError(f"Cannot build Excel file — the {name} sheet has no QnA pairs.")

    buffer = io.BytesIO()

    try:
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            for sheet_name in SHEET_ORDER:
                pairs = sheets[sheet_name]
                df = pd.DataFrame(
                    [(p["question"], p["answer"]) for p in pairs],
                    columns=["Questions", "Answers"],
                )
                df.to_excel(writer, sheet_name=sheet_name, index=False)

                worksheet = writer.sheets[sheet_name]
                header_fill = PatternFill(start_color=HEADER_COLOR, end_color=HEADER_COLOR, fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True)

                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(vertical="center")

                worksheet.column_dimensions["A"].width = 55
                worksheet.column_dimensions["B"].width = 75
                for row in worksheet.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    except Exception as e:
        raise ExcelBuildError(f"Failed to build Excel file: {e}")

    buffer.seek(0)
    return buffer.getvalue()
